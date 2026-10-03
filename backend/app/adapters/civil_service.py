"""Deterministic XLSX parser for official civil-service position workbooks."""
import io
import re
import zipfile
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from ..parsers import MAX_BYTES, ParseError
from .base import ParsedDocument, RawArtifact, SourceAdapter, ValidationResult

HEADER_ALIASES = {
    "position_code": ("职位代码", "职位编号", "岗位代码"),
    "department": ("部门名称", "部门", "招录机关"),
    "organization": ("用人司局", "机构名称", "招录单位", "单位名称"),
    "title": ("职位名称", "岗位名称", "招考职位"),
    "headcount": ("招考人数", "招录人数", "计划人数"),
    "education": ("学历", "学历要求"),
    "degree": ("学位", "学位要求"),
    "majors": ("专业", "专业要求"),
    "political_status": ("政治面貌",),
    "grassroots_years": ("基层工作最低年限", "基层工作年限"),
    "fresh_graduate": ("服务基层项目工作经历", "应届生要求", "应届要求"),
    "work_region": ("工作地点", "工作地区", "落户地点"),
    "remarks": ("备注", "其他条件", "说明"),
}
REQUIRED = {"position_code", "title", "education", "majors"}
RULE_FIELDS = ("education", "degree", "majors", "political_status", "grassroots_years", "fresh_graduate")


@dataclass
class CivilRecord:
    values: dict[str, Any]
    cells: dict[str, str]
    sheet: str
    row: int


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


class CivilServiceWorkbookAdapter(SourceAdapter):
    source_code = "CM-CS-001"

    def parse(self, raw: RawArtifact) -> ParsedDocument:
        if not raw.content or len(raw.content) > MAX_BYTES:
            raise ParseError("文件为空或超过 4 MiB")
        try:
            with zipfile.ZipFile(io.BytesIO(raw.content)) as archive:
                if len(archive.infolist()) > 2000 or any("vbaproject" in x.filename.lower() for x in archive.infolist()):
                    raise ParseError("工作簿包含宏或文件数量异常")
        except zipfile.BadZipFile as exc:
            raise ParseError("不是有效的 XLSX 工作簿") from exc
        book = load_workbook(io.BytesIO(raw.content), read_only=True, data_only=False, keep_links=False)
        records: list[dict[str, Any]] = []
        evidence: list[dict[str, str]] = []
        try:
            for sheet in book.worksheets:
                header_row, columns = self._find_header(sheet)
                if not header_row:
                    continue
                for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1), start=header_row + 1):
                    values, cells = {}, {}
                    for field, column in columns.items():
                        cell = row[column - 1]
                        if cell.data_type == "f":
                            raise ParseError(f"{sheet.title}!{cell.coordinate} 含公式，不能作为确定性事实")
                        values[field] = clean(cell.value)
                        cells[field] = f"sheet={sheet.title};cell={cell.coordinate}"
                    if not any(values.values()):
                        continue
                    if not values.get("position_code") and not values.get("title"):
                        continue
                    values["headcount"] = self._headcount(values.get("headcount"), sheet.title, row_number)
                    record = CivilRecord(values=values, cells=cells, sheet=sheet.title, row=row_number)
                    check = self.validate(record)
                    if not check.valid:
                        raise ParseError("; ".join(check.errors))
                    records.append({**values, "_cells": cells, "_sheet": sheet.title, "_row": row_number})
                    for field, value in values.items():
                        if field.startswith("_") or value in ("", None):
                            continue
                        evidence.append({"record_key": values["position_code"], "field": field,
                                         "evidence_location": cells[field],
                                         "quote_or_normalized_fact": f"{field}={value}", "extractor": "parser"})
        finally:
            book.close()
        if not records:
            raise ParseError("未找到包含职位代码、职位名称、学历和专业列的工作表")
        return ParsedDocument(records=records, evidence=evidence,
                              metadata={"source_item_id": raw.source_item_id, "canonical_url": raw.canonical_url})

    def _find_header(self, sheet):
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)), start=1):
            normalized = {clean(cell.value): index for index, cell in enumerate(row, start=1) if clean(cell.value)}
            columns = {}
            for field, aliases in HEADER_ALIASES.items():
                for label, index in normalized.items():
                    if label in aliases:
                        columns[field] = index
                        break
            if REQUIRED.issubset(columns):
                return row_number, columns
        return None, {}

    def _headcount(self, value, sheet, row):
        if value in ("", None):
            return None
        try:
            result = int(str(value).strip())
        except ValueError as exc:
            raise ParseError(f"{sheet} 第 {row} 行招录人数不是整数") from exc
        if result < 0 or result > 100000:
            raise ParseError(f"{sheet} 第 {row} 行招录人数超出合理范围")
        return result

    def normalize(self, parsed: ParsedDocument):
        return [{key: value for key, value in row.items() if not key.startswith("_")} for row in parsed.records]

    def extract_rules(self, parsed: ParsedDocument):
        rules = []
        evidence_index = {(e["record_key"], e["field"]): e for e in parsed.evidence}
        for row in parsed.records:
            for field in RULE_FIELDS:
                value = row.get(field)
                if not value:
                    continue
                rules.append({"record_key": row["position_code"], "field": field, "operator": "text_match",
                              "expected": str(value), "description": f"{field}: {value}", "confidence": 1.0,
                              "review_status": "pending", "extractor": "parser",
                              "evidence": evidence_index[(row["position_code"], field)]})
            if row.get("remarks"):
                rules.append({"record_key": row["position_code"], "field": "remarks", "operator": "candidate",
                              "expected": row["remarks"], "description": "备注中的复杂条件待人工确认", "confidence": 0.0,
                              "review_status": "pending", "extractor": "agent_candidate",
                              "evidence": evidence_index[(row["position_code"], "remarks")]})
        return rules

    def validate(self, record: CivilRecord) -> ValidationResult:
        missing = [field for field in REQUIRED if not record.values.get(field)]
        errors = [f"{record.sheet} 第 {record.row} 行缺少 {field}" for field in missing]
        return ValidationResult(valid=not errors, errors=errors)
