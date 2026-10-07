"""Bounded text extraction adapters. Never evaluate scripts, spreadsheet formulae or macros."""
import io
import os
import re
import zipfile
from dataclasses import dataclass, field

MAX_BYTES = 4 * 1024 * 1024
MAX_TEXT = 200000


@dataclass
class Parsed:
    text: str
    title: str
    format: str
    evidence: list = field(default_factory=list)


class ParseError(ValueError):
    pass


def extract(data: bytes, format: str) -> Parsed:
    if not data or len(data) > MAX_BYTES:
        raise ParseError("文件为空或超过 4 MiB")
    title = ""
    evidence = []
    if format == "html":
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(data, "html.parser")
        title = soup.title.get_text(" ", strip=True)[:160] if soup.title else ""
        for node in soup(["script", "style", "noscript", "nav", "footer", "header"]):
            node.decompose()
        main = soup.find("article") or soup.find("main") or soup.body or soup
        text = main.get_text("\n", strip=True)
    elif format == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseError("不支持加密 PDF")
        if len(reader.pages) > 100:
            raise ParseError("PDF 超过 100 页")
        parts, has_text, scanned = [], False, []
        for i, page in enumerate(reader.pages):
            # Some first-party PDF fonts decode invisible glyphs as NUL. PostgreSQL
            # rejects NUL in text fields; keep original bytes, sanitize only extracted text.
            page_text = (page.extract_text() or "").replace('\x00', '')
            if len(page_text.strip()) < 20:
                scanned.append(i+1)
            has_text = has_text or bool(page_text.strip())
            parts.append(f"[第 {i + 1} 页]\n" + page_text)
            if sum(map(len, parts)) > MAX_TEXT:
                raise ParseError("抽取文本超过 200000 字符")
        if scanned:
            from .ocr import recognize_pdf
            evidence = recognize_pdf(data, scanned)
            for number in scanned:
                lines=[p for p in evidence if p['page']==number]
                confidence=sum(p['confidence'] for p in lines)/len(lines)
                if confidence<65:
                    raise ParseError('OCR 置信度不足 65%，需人工核对')
                parts[number-1]=f'[第 {number} 页]\n'+'\n'.join(p['text'] for p in lines)
        text = "\n\n".join(parts)
    elif format == 'docx':
        from xml.etree import ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist()) > 1000 or sum(i.file_size for i in z.infolist()) > 20*1024*1024:
                raise ParseError('Word 解压体积或文件数量过大')
            if any('vbaproject' in i.filename.lower() for i in z.infolist()):
                raise ParseError('不支持含宏的 Word 文档')
            xml=z.read('word/document.xml')
        if re.search(br'<!\s*(?:DOCTYPE|ENTITY)',xml.replace(b'\x00',b''),re.I):
            raise ParseError('Word 正文包含不支持的 XML 实体')
        root=ET.fromstring(xml)
        ns='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        parts=[]
        for index, paragraph in enumerate(root.iter(ns+'p'),1):
            if index>5000:
                raise ParseError('Word 超过 5000 个段落')
            value=''.join(n.text or '' for n in paragraph.iter(ns+'t')).strip()
            if value:parts.append(f'[段落 {index}] '+value)
        text='\n'.join(parts)
    elif format in {"xlsx", "xlsx_attachment"}:
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist()) > 1000 or sum(i.file_size for i in z.infolist()) > 20 * 1024 * 1024:
                raise ParseError("Excel 解压体积或文件数量过大")
            if any("vbaproject" in i.filename.lower() for i in z.infolist()):
                raise ParseError("不支持含宏的工作簿")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) > 20:
                raise ParseError("Excel 超过 20 个工作表")
            parts, total, rows = [], 0, 0
            for sheet in book.worksheets:
                parts.append(f"[工作表：{sheet.title}]")
                for row_index,row in enumerate(sheet.iter_rows(),1):
                    rows += 1
                    if rows > 5000 or len(row) > 100:
                        raise ParseError("Excel 超过 5000 行或 100 列")
                    # Never use a stale cached formula value as a policy fact.
                    values = ["[公式，需人工核对]" if c.data_type == "f" else str(c.value) if c.value is not None else "" for c in row]
                    line = " | ".join(values).strip(" |")
                    if line:
                        total += len(line)
                        if total > MAX_TEXT:
                            raise ParseError("抽取文本超过 200000 字符")
                        parts.append((f'[行 {row_index}] ' if format=='xlsx_attachment' else '')+line)
            text = "\n".join(parts)
        finally:
            book.close()
    elif format in {'doc','xls'}:
        from .legacy_office import word, spreadsheet
        text=word(data) if format=='doc' else spreadsheet(data)
    elif format == 'image':
        from .ocr import recognize_bytes
        evidence=recognize_bytes(data)
        if sum(p['confidence'] for p in evidence)/len(evidence)<65:
            raise ParseError('OCR 置信度不足 65%，需人工核对')
        text='[第 1 页]\n'+'\n'.join(p['text'] for p in evidence)
    elif format == "text":
        text = data.decode("utf-8-sig")
    else:
        raise ParseError("只支持 HTML、PDF、DOC/DOCX、XLS/XLSX、图片和 UTF-8 文本")
    text = re.sub(r"\n{3,}", "\n\n", text.replace('\x00','')).strip()
    if len(text) < 20:
        raise ParseError("有效正文不足 20 字符，请人工核对")
    if len(text) > MAX_TEXT:
        raise ParseError("抽取文本超过 200000 字符")
    return Parsed(text=text, title=title, format=format, evidence=evidence)


def extract_isolated(data: bytes, format: str) -> Parsed:
    """Keep parser failures/timeouts out of the API/worker process."""
    import json
    import subprocess
    import sys
    if not data or len(data) > MAX_BYTES:
        raise ParseError("文件为空或超过 4 MiB")
    try:
        # The API is commonly launched from the repository root, while the isolated
        # module lives under ``backend/app``. Pass an explicit import root so the
        # parser works identically under pytest, uvicorn, workers and packaged runs.
        backend_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        env = os.environ.copy()
        env["PYTHONPATH"] = backend_root + os.pathsep + env.get("PYTHONPATH", "")
        result = subprocess.run([sys.executable, "-m", "app.parsers", format], input=data,
                                capture_output=True, timeout=120 if format=='pdf' else 45, check=False, cwd=backend_root, env=env)
    except subprocess.TimeoutExpired as exc:
        raise ParseError("解析超时，请人工处理") from exc
    if result.returncode or len(result.stdout) > MAX_TEXT * 6 + 2048:
        raise ParseError("文件无法解析或超出资源限制，请人工处理")
    body = json.loads(result.stdout)
    if "error" in body:
        raise ParseError(body["error"])
    return Parsed(**body)


if __name__ == "__main__":
    import json
    import sys
    from dataclasses import asdict
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_CPU, (90, 90))
    except (ImportError, OSError, ValueError):
        pass  # Wall time and input/output bounds still apply on non-POSIX hosts.
    try:
        print(json.dumps(asdict(extract(sys.stdin.buffer.read(MAX_BYTES + 1), sys.argv[1])), ensure_ascii=False))
    except Exception as error:  # noqa: BLE001 - isolate malformed third-party parser inputs
        print(json.dumps({"error": str(error) if isinstance(error, ValueError) else "文件格式无效，无法提取正文"}))
